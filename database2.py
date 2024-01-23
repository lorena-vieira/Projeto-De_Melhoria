import tkinter as tk
from tkinter import messagebox
from docxtpl import DocxTemplate
from pymongo import MongoClient
from tkinter import filedialog
from docx import Document
import Tela1

root = tk.Tk()
app = Tela1.DataEntryApp(root)
user_input = app.submit_data()
root.mainloop()

def generate_diagnostic_report(output_path, user_input):
    # Carregar o modelo docx
    doc = DocxTemplate(r'C:\Users\Lorena\Documents\PONTES-MH\ProjetoDeMelhoria\arquivoBasePython\RelatorioMetalica.docx')

    # Adicionar os dados do usuário ao contexto
    context = user_input

    # Renderizar o documento
    doc.render(context)

    # Salvar o relatório gerado
    doc.save(output_path)

# Exemplo de uso
if __name__ == '__main__':
    output_path = 'diagnostic_report.docx'

    # Obter informações do usuário
    user_input = Tela1.submit_data()

    # Gerar relatório
    generate_diagnostic_report(output_path, user_input)

    print(f"Relatório gerado com sucesso em {output_path}")


#A seguir a criação de tela para seleção de patologias, inserção de fotos e classificação (notas)


CONNECTION_STRING = "mongodb+srv://lorenavieira:lorenas2fabinho@cluster0.f65bx3a.mongodb.net/"
client = MongoClient(CONNECTION_STRING)
db = client['PatologiasDescricoes']
collection = db['anomalias']

class DiagnosticsApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Aplicação de Diagnóstico")

                # Interface do Usuário
        self.anomaly_label = tk.Label(root, text="Selecione a Anomalia:")
        self.anomaly_label.pack()

        # Obtém anomalias do banco de dados
        self.anomalies = self.fetch_anomalies()
        self.selected_anomalies = []

        # Checkbox para anomalias
        self.anomalies_frame = tk.Frame(root)
        self.anomalies_frame.pack(pady=10)
        self.anomalies_label = tk.Label(self.anomalies_frame, text="Escolha as Anomalias:")
        self.anomalies_label.pack()

        for anomaly in self.anomalies:
            checkbox = tk.Checkbutton(self.anomalies_frame, text=anomaly, variable=anomaly)
            checkbox.pack()

        # Botão para adicionar fotos e notas
        self.add_photos_button = tk.Button(root, text="Adicionar Fotos e Notas", command=self.add_photos_and_notes)
        self.add_photos_button.pack(pady=10)
        self.anomaly_var = tk.StringVar(root)
        self.anomaly_var.set(self.anomalies[0])  # Define o valor padrão

        self.anomaly_menu = tk.OptionMenu(root, self.anomaly_var, *self.anomalies)
        self.anomaly_menu.pack()

        self.photos_entry_label = tk.Label(root, text="Insira aqui quantidade de fotos da anomalia:")
        self.photos_entry_label.pack()

        self.photos_entry = tk.Entry(root)
        self.photos_entry.pack()

        self.durability_note_label = tk.Label(root, text="Nota para Durabilidade (1-5):")
        self.durability_note_label.pack()

        self.durability_note_entry = tk.Entry(root)
        self.durability_note_entry.pack()

        self.structural_note_label = tk.Label(root, text="Nota Estrutural (1-5):")
        self.structural_note_label.pack()

        self.structural_note_entry = tk.Entry(root)
        self.structural_note_entry.pack()

        self.submit_button = tk.Button(root, text="Submeter", command=self.generate_report)
        self.submit_button.pack()

    def fetch_anomalies(self):
        anomalies = collection.find({}, {"anomaly_name": 1})
        return [anomaly for anomaly in anomalies]

    def generate_report(self):
        anomaly = self.anomaly_var.get()
        photos_count = self.photos_entry.get()
        durability_note = self.durability_note_entry.get()
        structural_note = self.structural_note_entry.get()

        # Validar entrada (adapte conforme necessário)
        if not photos_count.isdigit() or int(photos_count) % 2 != 0:
            messagebox.showerror("Erro", "Informe uma quantidade par de fotos.")
            return

        try:
            durability_note = int(durability_note)
            structural_note = int(structural_note)
            if not (1 <= durability_note <= 5) or not (1 <= structural_note <= 5):
                raise ValueError("Notas devem estar entre 1 e 5.")
        except ValueError as e:
            messagebox.showerror("Erro", str(e))
            return


    def add_photos_and_notes(self):
        selected_anomalies = [anomaly for anomaly in self.anomalies if anomaly is not None]

        for anomaly in selected_anomalies:
            # Adicione aqui a lógica para adicionar campos de fotos e notas
            print(f"Anomalia: {anomaly}")

            # Exemplo: Pedir ao usuário para selecionar uma foto
            file_path = filedialog.askopenfilename(title=f"Selecione uma foto para {anomaly}", filetypes=[("JPEG files", "*.jpg;*.jpeg")])
            print(f"Foto selecionada: {file_path}")

            # Exemplo: Pedir ao usuário para inserir notas de durabilidade e estrutural
            durabilidade = input(f"Insira a nota de durabilidade para {anomaly} (1-5): ")
            estrutural = input(f"Insira a nota estrutural para {anomaly} (1-5): ")
            print(f"Notas - Durabilidade: {durabilidade}, Estrutural: {estrutural}")


        # Botão para aplicar alterações ao documento .docx
        self.apply_changes_button = tk.Button(root, text="Aplicar Alterações ao Documento .docx", command=self.apply_changes_to_docx)
        self.apply_changes_button.pack(pady=10)

        # Armazenar dados do usuário
        self.user_data["anomalies"].append({
                "anomaly": anomaly,
                "photo": file_path,
                "durabilidade": durabilidade,
                "estrutural": estrutural
            })

    def apply_changes_to_docx(self):
        # Carregar o documento .docx existente
        doc = Document(r'C:\Users\Lorena\PycharmProjects\Projeto-de-melhoria\diagnostic_report.docx')

        # Iterar sobre os parágrafos do documento
        for paragraph in doc.paragraphs:
            # Substituir placeholders nos parágrafos
            for anomaly_data in self.user_data["anomalies"]:
                placeholder = f"{{{anomaly_data['anomaly']}}}"
                if placeholder in paragraph.text:
                    # Substituir o placeholder com os dados do usuário
                    paragraph.text = paragraph.text.replace(placeholder, f"Anomalia: {anomaly_data['anomaly']}, Durabilidade: {anomaly_data['durabilidade']}, Estrutural: {anomaly_data['estrutural']}")

        # Salvar o documento atualizado
        doc.save("RelatorioMetalica_atualizado.docx")


    def generate_diagnostic_report(self, anomaly, photos_count, durability_note, structural_note):
        # Carregar o modelo docx
        doc = DocxTemplate(r'C:\Users\Lorena\PycharmProjects\Projeto-de-melhoria\diagnostic_report.docx')

        # Adicionar os dados ao contexto
        context = {
            'anomaly': anomaly,
            'photos_count': photos_count,
            'durability_note': durability_note,
            'structural_note': structural_note,
        }

        # Renderizar o documento
        doc.render(context)

        # Salvar o relatório gerado (substitua 'output_report.docx' pelo nome desejado)
        doc.save('output_report.docx')

        messagebox.showinfo("Relatório Gerado", "Relatório gerado com sucesso.")


if __name__ == '__main__':
    root = tk.Tk()
    app = DiagnosticsApp(root)

    label = tk.Label(root, text="Olá, tkinter!")
    label.pack()



    root.mainloop()





